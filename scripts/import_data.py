#!/usr/bin/env python3
"""Import vinyl spreadsheet and Notion media tracker into MediaShelf."""

import csv
import json
import sys
import urllib.request

import openpyxl

API_BASE = "http://localhost:8000/api/v1"

NOTION_TYPE_MAP = {
    "TV Series": "show",
    "Movie": "movie",
    "Book": "book",
    "Documentary": "documentary",
    "Audiobook": "audiobook",
}

NOTION_STATUS_MAP = {
    "Completed": "completed",
    "In Progress": "in_progress",
    "Dropped": "dropped",
    "Want to Watch/Read": "wishlist",
}


def get_token(username: str, password: str) -> str:
    data = json.dumps({"username": username, "password": password}).encode()
    req = urllib.request.Request(
        f"{API_BASE}/auth/login",
        data=data,
        headers={"Content-Type": "application/json"},
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read())["access_token"]


def add_item(token: str, item: dict) -> tuple[bool, str]:
    data = json.dumps(item).encode()
    req = urllib.request.Request(
        f"{API_BASE}/collection/",
        data=data,
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {token}",
        },
    )
    try:
        with urllib.request.urlopen(req) as resp:
            return True, "added"
    except urllib.error.HTTPError as e:
        if e.code == 409:
            return False, "duplicate"
        body = e.read().decode()
        return False, f"error {e.code}: {body}"


def count_stars(rating_str: str) -> int | None:
    if not rating_str:
        return None
    count = rating_str.count("⭐")
    return count if 1 <= count <= 5 else None


def import_vinyl(token: str, xlsx_path: str):
    print("\n=== Importing Vinyl (Owned) ===")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    # Owned sheet
    ws = wb["Owned"]
    added, dupes, errors = 0, 0, 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        artist, genre, album, year, decade, color, instrumental, apple_music = row[:8]
        if not artist or not album:
            continue

        item = {
            "media_type": "vinyl",
            "title": str(album).strip(),
            "status": "owned",
            "metadata": {
                "artist": str(artist).strip(),
                "genre": str(genre).strip() if genre else None,
                "year": int(year) if year else None,
                "decade": str(decade).strip() if decade else None,
                "vinyl_color": str(color).strip() if color else None,
                "instrumental": str(instrumental).strip() == "Yes" if instrumental else False,
            },
            "tags": [g.strip().lower() for g in str(genre).split(",") if g.strip()] if genre else [],
        }
        # Remove None values from metadata
        item["metadata"] = {k: v for k, v in item["metadata"].items() if v is not None}

        ok, msg = add_item(token, item)
        if ok:
            added += 1
        elif msg == "duplicate":
            dupes += 1
        else:
            errors += 1
            print(f"  Error: {album} by {artist} - {msg}")

    print(f"  Owned: {added} added, {dupes} duplicates, {errors} errors")

    # Wishlist sheet
    ws2 = wb["Wishlist"]
    added2, dupes2, errors2 = 0, 0, 0
    for row in ws2.iter_rows(min_row=2, values_only=True):
        title, artist = row[0], row[1] if len(row) > 1 else None
        if not title:
            continue

        item = {
            "media_type": "vinyl",
            "title": str(title).strip(),
            "status": "wishlist",
            "metadata": {"artist": str(artist).strip()} if artist else {},
        }

        ok, msg = add_item(token, item)
        if ok:
            added2 += 1
        elif msg == "duplicate":
            dupes2 += 1
        else:
            errors2 += 1
            print(f"  Error: {title} - {msg}")

    print(f"  Wishlist: {added2} added, {dupes2} duplicates, {errors2} errors")


def import_notion(token: str, csv_path: str):
    print("\n=== Importing Notion Media Tracker ===")
    added, dupes, errors, skipped = 0, 0, 0, 0

    with open(csv_path, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            title = row.get("Title", "").strip()
            if not title:
                continue

            notion_type = row.get("Type", "").strip()
            media_type = NOTION_TYPE_MAP.get(notion_type)
            if not media_type:
                print(f"  Skipping unknown type '{notion_type}': {title}")
                skipped += 1
                continue

            status = NOTION_STATUS_MAP.get(row.get("Status", "").strip(), "owned")
            rating = count_stars(row.get("Rating", ""))
            genre = row.get("Genre", "").strip()

            item = {
                "media_type": media_type,
                "title": title,
                "status": status,
                "rating": rating,
                "metadata": {},
                "tags": [g.strip().lower() for g in genre.split(",") if g.strip()] if genre else [],
            }
            if genre:
                item["metadata"]["genre"] = genre

            ok, msg = add_item(token, item)
            if ok:
                added += 1
            elif msg == "duplicate":
                dupes += 1
            else:
                errors += 1
                print(f"  Error: {title} - {msg}")

    print(f"  {added} added, {dupes} duplicates, {errors} errors, {skipped} skipped")


def main():
    username = input("Username: ") if len(sys.argv) < 2 else sys.argv[1]
    password = input("Password: ") if len(sys.argv) < 3 else sys.argv[2]

    print(f"Logging in as {username}...")
    token = get_token(username, password)
    print("Authenticated.")

    xlsx_path = "/Users/richardmarchetti/Downloads/6J Vinyl.xlsx"
    csv_path = "/tmp/notion_data/Private & Shared/Media Tracker 1e49e5e56ff3802f944eee44fcd78d97_all.csv"

    import_vinyl(token, xlsx_path)
    import_notion(token, csv_path)

    print("\nDone!")


if __name__ == "__main__":
    main()
